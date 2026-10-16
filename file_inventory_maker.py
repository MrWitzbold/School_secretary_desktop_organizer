#!/usr/bin/env python3
"""
Recursive Folder Inventory (com OCR opcional)
---------------------------------------------
Cria folder_inventory.txt contendo:
- Estrutura recursiva completa de pastas e arquivos
- Caminhos relativos, extensões e tamanhos
- Até 5 linhas de conteúdo por arquivo
- OCR pt-BR para imagens e PDFs escaneados (via Tesseract)
- Erros e observações para arquivos não suportados/ilegíveis

Dependências Python (instale só o que precisar):
  pip install pymupdf python-docx openpyxl xlrd python-pptx odfpy
  pip install pytesseract Pillow

OCR também requer o binário Tesseract com pacote de português:
  Windows: https://github.com/UB-Mannheim/tesseract/wiki
  Linux:   sudo apt install tesseract-ocr tesseract-ocr-por
  macOS:   brew install tesseract tesseract-lang

Coloque este script na pasta que deseja escanear.
Execute: python file_inventory_maker.py
"""

from pathlib import Path
from datetime import datetime
import os
import sys
import csv
import json
import zipfile
import mimetypes
import subprocess
import shutil
import io

ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "folder_inventory.txt"
MAX_LINES = 5
MAX_FILE_SIZE_MB = 50

# Quantas páginas de PDF tentar via OCR se não houver texto extraível
PDF_OCR_MAX_PAGES = 5
# DPI para rasterizar PDFs antes do OCR (200 costuma ser um bom equilíbrio)
PDF_OCR_DPI = 200

TEXT_EXTENSIONS = {
    ".txt", ".md", ".rst", ".log", ".ini", ".cfg", ".conf",
    ".toml", ".yaml", ".yml", ".xml", ".html", ".htm", ".css",
    ".js", ".jsx", ".ts", ".tsx", ".py", ".pyw", ".java",
    ".c", ".h", ".cpp", ".hpp", ".cs", ".go", ".rs", ".php",
    ".rb", ".sh", ".bash", ".zsh", ".ps1", ".bat", ".cmd",
    ".sql", ".r", ".m", ".mat", ".tex", ".bib", ".diff",
    ".patch", ".env", ".gitignore", ".dockerfile", ".makefile",
    ".properties", ".gradle", ".kt", ".swift", ".dart", ".vue",
    ".svelte", ".ipynb", ".graphql", ".proto", ".asm",
}

IMAGE_EXTENSIONS = {
    ".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp",
    ".tif", ".tiff", ".ico", ".heic", ".svg"
}

# Extensões que tentamos OCR diretamente
OCR_IMAGE_EXTENSIONS = {
    ".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp",
    ".tif", ".tiff", ".ico"
}

AUDIO_VIDEO_EXTENSIONS = {
    ".mp3", ".wav", ".ogg", ".flac", ".m4a", ".aac",
    ".mp4", ".mkv", ".avi", ".mov", ".webm", ".wmv"
}

ARCHIVE_EXTENSIONS = {
    ".zip", ".rar", ".7z", ".tar", ".gz", ".bz2", ".xz"
}


# ---------------------------------------------------------------------------
# Utilitários de texto
# ---------------------------------------------------------------------------

def clean_text(value):
    value = str(value).replace("\x00", "")
    return value.replace("\r\n", "\n").replace("\r", "\n")


def first_lines(text, limit=MAX_LINES):
    text = clean_text(text)
    lines = text.split("\n")

    while lines and not lines[-1].strip():
        lines.pop()

    if not lines or not text.strip():
        return "[No readable text content found]"

    result = lines[:limit]

    if len(lines) > limit:
        result.append("[Preview truncated; more content exists]")

    return "\n".join(result)


# ---------------------------------------------------------------------------
# OCR
# ---------------------------------------------------------------------------

def run_ocr(pil_image):
    """Executa Tesseract em uma imagem PIL. Retorna (texto, erro)."""
    try:
        import pytesseract
    except ImportError:
        return None, "Install pytesseract and Pillow: pip install pytesseract Pillow"

    # Tenta pt-BR primeiro; cai para inglês/default se faltar o pacote.
    for lang in ("por+eng", "por", "eng"):
        try:
            text = pytesseract.image_to_string(pil_image, lang=lang)
            if text and text.strip():
                return text, None
        except pytesseract.TesseractNotFoundError:
            return None, (
                "Tesseract binary not found. Install Tesseract OCR "
                "(with Portuguese language pack) and add it to PATH."
            )
        except pytesseract.TesseractError:
            # Normalmente idioma não instalado — tenta o próximo.
            continue
        except Exception as exc:
            return None, f"OCR error: {exc}"

    return None, "OCR found no text"


def ocr_image_file(path):
    """Abre uma imagem com Pillow e roda OCR."""
    try:
        from PIL import Image
    except ImportError:
        return None, "Install Pillow: pip install Pillow"

    try:
        with Image.open(path) as img:
            # Converte para RGB (Tesseract prefere)
            if img.mode not in ("RGB", "L"):
                img = img.convert("RGB")

            # Upscale se for pequena — ajuda bastante no OCR
            if img.width < 1200:
                scale = 1200 / img.width
                img = img.resize(
                    (int(img.width * scale), int(img.height * scale))
                )

            return run_ocr(img)
    except Exception as exc:
        return None, f"Image error: {exc}"


def ocr_pdf(path, max_pages=PDF_OCR_MAX_PAGES, dpi=PDF_OCR_DPI):
    """Rasteriza páginas do PDF com pymupdf e roda OCR."""
    try:
        import fitz  # pymupdf
    except ImportError:
        return None, "Install pymupdf: pip install pymupdf"

    try:
        from PIL import Image
    except ImportError:
        return None, "Install Pillow: pip install Pillow"

    try:
        doc = fitz.open(str(path))
    except Exception as exc:
        return None, f"PDF OCR error: {exc}"

    parts = []
    try:
        zoom = dpi / 72.0
        matrix = fitz.Matrix(zoom, zoom)

        for page_index in range(min(len(doc), max_pages)):
            page = doc[page_index]
            pix = page.get_pixmap(matrix=matrix, alpha=False)
            img = Image.open(io.BytesIO(pix.tobytes("png")))

            text, err = run_ocr(img)

            if err and "Install" in err:
                return None, err

            if text and text.strip():
                parts.append(f"[Page {page_index + 1} — OCR]\n{text}")

        if not parts:
            return None, "OCR found no text in scanned PDF"

        if len(doc) > max_pages:
            parts.append(
                f"[OCR stopped after {max_pages} pages]"
            )

        return "\n".join(parts), None
    except Exception as exc:
        return None, f"PDF OCR error: {exc}"
    finally:
        try:
            doc.close()
        except Exception:
            pass


# ---------------------------------------------------------------------------
# Extratores por formato
# ---------------------------------------------------------------------------

def read_text_file(path):
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            with path.open("r", encoding=encoding) as file:
                return file.read(), None
        except UnicodeDecodeError:
            continue

    return None, "Could not decode text"


def extract_pdf(path):
    """Extrai texto com pymupdf; cai para OCR se for escaneado."""
    try:
        import fitz  # pymupdf
    except ImportError:
        return None, "Install pymupdf: pip install pymupdf"

    try:
        doc = fitz.open(str(path))

        parts = []
        total_chars = 0

        for page_number, page in enumerate(doc):
            text = page.get_text() or ""
            if text.strip():
                parts.append(f"[Page {page_number + 1}]\n{text}")
                total_chars += len(text)

            if total_chars > 100_000:
                break

        doc.close()

        if parts:
            return "\n".join(parts), None

        # Sem texto: tenta OCR
        return ocr_pdf(path)

    except Exception as exc:
        return None, f"PDF error: {exc}"


def extract_docx(path):
    try:
        from docx import Document
    except ImportError:
        return None, "Install python-docx: pip install python-docx"

    try:
        document = Document(str(path))
        parts = []

        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                parts.append(paragraph.text)

        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))

        if not parts:
            return None, "DOCX contained no text"

        return "\n".join(parts), None

    except Exception as exc:
        return None, f"DOCX error: {exc}"


def extract_xlsx(path):
    try:
        from openpyxl import load_workbook
    except ImportError:
        return None, "Install openpyxl: pip install openpyxl"

    try:
        workbook = load_workbook(
            filename=str(path), read_only=True, data_only=True
        )
        parts = []

        for sheet in workbook.worksheets:
            parts.append(f"[Sheet: {sheet.title}]")

            for row in sheet.iter_rows(values_only=True):
                values = [
                    str(value) if value is not None else "" for value in row
                ]
                if any(v for v in values if v):
                    parts.append(" | ".join(values))

                if len(parts) > 1000:
                    parts.append("[Spreadsheet preview limit reached]")
                    break

            if len(parts) > 1000:
                break

        workbook.close()

        if not parts:
            return None, "Workbook is empty"

        return "\n".join(parts), None

    except Exception as exc:
        return None, f"XLSX error: {exc}"


def extract_xls(path):
    try:
        import xlrd
    except ImportError:
        return None, "Install xlrd: pip install xlrd"

    try:
        workbook = xlrd.open_workbook(str(path), on_demand=True)
        parts = []

        for sheet in workbook.sheets():
            parts.append(f"[Sheet: {sheet.name}]")

            for row_index in range(min(sheet.nrows, 1000)):
                values = [str(v) for v in sheet.row_values(row_index)]
                if any(v.strip() for v in values):
                    parts.append(" | ".join(values))

        workbook.release_resources()

        if not parts:
            return None, "Workbook is empty"

        return "\n".join(parts), None

    except Exception as exc:
        return None, f"XLS error: {exc}"


def extract_pptx(path):
    try:
        from pptx import Presentation
    except ImportError:
        return None, "Install python-pptx: pip install python-pptx"

    try:
        presentation = Presentation(str(path))
        parts = []

        for number, slide in enumerate(presentation.slides, start=1):
            parts.append(f"[Slide {number}]")

            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    if shape.text.strip():
                        parts.append(shape.text)

                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        parts.append(
                            " | ".join(cell.text for cell in row.cells)
                        )

        if not parts:
            return None, "Presentation has no text"

        return "\n".join(parts), None

    except Exception as exc:
        return None, f"PPTX error: {exc}"


def extract_odt(path):
    try:
        from odf.opendocument import load
        from odf import text as odf_text
        from odf import teletype
    except ImportError:
        return None, "Install odfpy: pip install odfpy"

    try:
        document = load(str(path))
        paragraphs = document.getElementsByType(odf_text.P)
        parts = [teletype.extractText(p) for p in paragraphs]
        if not parts:
            return None, "ODT has no text"
        return "\n".join(parts), None

    except Exception as exc:
        return None, f"ODT error: {exc}"


def extract_odp(path):
    try:
        from odf.opendocument import load
        from odf import draw, text as odf_text, teletype
    except ImportError:
        return None, "Install odfpy: pip install odfpy"

    try:
        document = load(str(path))
        parts = []

        for slide in document.getElementsByType(draw.Page):
            parts.append("[Slide]")
            for paragraph in slide.getElementsByType(odf_text.P):
                parts.append(teletype.extractText(paragraph))

        if not parts:
            return None, "ODP has no text"
        return "\n".join(parts), None

    except Exception as exc:
        return None, f"ODP error: {exc}"


def extract_doc(path):
    """DOC legado via antiword, se disponível."""
    executable = shutil.which("antiword")

    if executable:
        try:
            result = subprocess.run(
                [executable, str(path)],
                capture_output=True,
                text=True,
                timeout=20,
                errors="replace",
            )
            if result.returncode == 0:
                return result.stdout, None
            return None, "antiword could not read this DOC file"
        except Exception as exc:
            return None, f"DOC error: {exc}"

    return None, "Legacy DOC: install antiword or convert to DOCX"


def extract_ods(path):
    try:
        from xml.etree import ElementTree as ET

        with zipfile.ZipFile(path) as archive:
            xml = archive.read("content.xml")

        root = ET.fromstring(xml)
        rows = []
        for element in root.iter():
            if element.tag.endswith("}table-row"):
                cells = "".join(element.itertext())
                if cells.strip():
                    rows.append(cells)

        if not rows:
            return None, "ODS has no text"
        return "\n".join(rows), None

    except Exception as exc:
        return None, f"ODS error: {exc}"


# ---------------------------------------------------------------------------
# Roteador
# ---------------------------------------------------------------------------

def extract_content(path):
    ext = path.suffix.lower()
    name = path.name.lower()

    if ext in TEXT_EXTENSIONS or name in {
        "dockerfile", "makefile", ".gitignore", ".env"
    }:
        return read_text_file(path)

    if ext == ".csv":
        try:
            with path.open(
                "r", encoding="utf-8-sig", errors="replace", newline=""
            ) as file:
                rows = list(csv.reader(file))
            if not rows:
                return None, "Empty CSV"
            return "\n".join(" | ".join(r) for r in rows[:1000]), None
        except Exception as exc:
            return None, f"CSV error: {exc}"

    if ext in {".json", ".jsonl", ".ndjson"}:
        return read_text_file(path)

    if ext == ".pdf":
        return extract_pdf(path)

    if ext == ".docx":
        return extract_docx(path)

    if ext == ".doc":
        return extract_doc(path)

    if ext in {".xlsx", ".xlsm"}:
        return extract_xlsx(path)

    if ext == ".xls":
        return extract_xls(path)

    if ext == ".pptx":
        return extract_pptx(path)

    if ext == ".odt":
        return extract_odt(path)

    if ext == ".odp":
        return extract_odp(path)

    if ext == ".ods":
        return extract_ods(path)

    if ext in IMAGE_EXTENSIONS:
        if ext == ".svg":
            return read_text_file(path)

        if ext == ".heic":
            return None, "HEIC image: install pillow-heif to enable OCR"

        if ext in OCR_IMAGE_EXTENSIONS:
            return ocr_image_file(path)

        return None, "Image file; OCR not attempted"

    if ext in AUDIO_VIDEO_EXTENSIONS:
        return None, "Audio/video file; transcription not enabled"

    if ext in ARCHIVE_EXTENSIONS:
        if ext == ".zip":
            try:
                with zipfile.ZipFile(path) as archive:
                    names = archive.namelist()
                return (
                    "[Archive entries]\n" + "\n".join(names[:100]),
                    None,
                )
            except Exception as exc:
                return None, f"Archive error: {exc}"

        return None, "Archive detected; contents not extracted"

    # Última tentativa: arquivo parece texto?
    try:
        with path.open("rb") as file:
            sample = file.read(4096)

        if b"\x00" not in sample:
            return read_text_file(path)

        return None, "Binary or unsupported file format"

    except Exception as exc:
        return None, f"Could not inspect file: {exc}"


# ---------------------------------------------------------------------------
# Utilitários de arquivo
# ---------------------------------------------------------------------------

def safe_size(path):
    try:
        return path.stat().st_size
    except OSError:
        return 0


def format_size(size):
    if size < 1024:
        return f"{size} B"
    if size < 1024 ** 2:
        return f"{size / 1024:.1f} KB"
    if size < 1024 ** 3:
        return f"{size / (1024 ** 2):.2f} MB"
    return f"{size / (1024 ** 3):.2f} GB"


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    print(f"Scanning: {ROOT}")
    print("This may take a while for large folders...\n")

    directories = []
    files = []
    errors = []

    for current, dirnames, filenames in os.walk(ROOT):
        current_path = Path(current)

        dirnames[:] = sorted(
            d for d in dirnames
            if not (current_path / d).is_symlink()
        )

        for dirname in sorted(dirnames):
            directories.append(current_path / dirname)

        for filename in sorted(filenames):
            path = current_path / filename

            if path.resolve() == OUTPUT.resolve():
                continue

            if path.is_symlink():
                files.append((path, "SYMLINK"))
            else:
                files.append((path, None))

    total = len(files)
    report = []

    report.append("=" * 78)
    report.append("RECURSIVE FOLDER INVENTORY")
    report.append("=" * 78)
    report.append(f"Root folder: {ROOT}")
    report.append(f"Generated: {datetime.now().astimezone().isoformat()}")
    report.append(f"Folders found: {len(directories)}")
    report.append(f"Files found: {total}")
    report.append(f"Text preview limit: {MAX_LINES} lines per file")
    report.append("=" * 78)
    report.append("")
    report.append("FOLDER TREE")
    report.append("-" * 78)
    report.append(ROOT.name + "/")

    all_paths = sorted(
        directories + [p for p, _ in files],
        key=lambda p: (str(p.parent).lower(), p.name.lower()),
    )

    children = {}
    for path in all_paths:
        children.setdefault(path.parent, []).append(path)

    def add_tree(directory, prefix=""):
        entries = sorted(
            children.get(directory, []),
            key=lambda p: (not p.is_dir(), p.name.lower()),
        )

        for index, path in enumerate(entries):
            is_last = index == len(entries) - 1
            connector = "└── " if is_last else "├── "
            suffix = "/" if path.is_dir() else ""

            report.append(prefix + connector + path.name + suffix)

            if path.is_dir():
                add_tree(path, prefix + ("    " if is_last else "│   "))

    add_tree(ROOT)

    report.append("")
    report.append("=" * 78)
    report.append("FILE DETAILS AND CONTENT PREVIEWS")
    report.append("=" * 78)

    for index, (path, special) in enumerate(files, start=1):
        relative = path.relative_to(ROOT)
        ext = path.suffix.lower() or "[no extension]"
        size = safe_size(path)

        report.append("")
        report.append("#" * 78)
        report.append(f"FILE {index}/{total}")
        report.append(f"Name: {path.name}")
        report.append(f"Relative path: {relative}")
        report.append(f"Extension: {ext}")
        report.append(f"Size: {format_size(size)}")
        report.append(
            f"Type: {mimetypes.guess_type(path.name)[0] or 'unknown'}"
        )

        if special == "SYMLINK":
            report.append("Status: Symbolic link; not followed")
            continue

        if size > MAX_FILE_SIZE_MB * 1024 * 1024:
            report.append(
                f"Status: Skipped; larger than {MAX_FILE_SIZE_MB} MB"
            )
            continue

        try:
            content, error = extract_content(path)

            if error:
                report.append(f"Note: {error}")

            if content is not None:
                report.append("")
                report.append(f"FIRST {MAX_LINES} LINES OF CONTENT")
                report.append("-" * 40)
                report.append(first_lines(content))

        except Exception as exc:
            report.append(f"ERROR: {type(exc).__name__}: {exc}")
            errors.append((str(relative), str(exc)))

        if index % 25 == 0 or index == total:
            print(f"Processed {index}/{total} files...")

    report.append("")
    report.append("=" * 78)
    report.append("END OF INVENTORY")
    report.append(f"Root folder: {ROOT}")
    report.append(f"Folders: {len(directories)}")
    report.append(f"Files: {total}")
    report.append(f"Unexpected errors: {len(errors)}")
    report.append("=" * 78)

    try:
        OUTPUT.write_text("\n".join(report), encoding="utf-8")
    except OSError as exc:
        print(f"\nCould not write report: {exc}", file=sys.stderr)
        sys.exit(1)

    print("\nDone!")
    print(f"Report saved to: {OUTPUT}")
    print(f"Folders: {len(directories)}")
    print(f"Files: {total}")
    print(f"Report size: {format_size(safe_size(OUTPUT))}")
    print("\nSend folder_inventory.txt here when you want me to inspect it.")


if __name__ == "__main__":
    main()